"""Build the editable thesis from the reviewed Word template and the PDF's resolved source.

Run with the configured document-authoring Python runtime (python-docx and lxml).
Microsoft Office supplies MML2OMML.XSL; the stylesheet is not redistributed.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_SHA = "d6f5cef5d56bdcb7f6c78989d68f8c586ba9b9289ecf5b765ebd34a8e26d3d59"
REFERENCE = re.compile(r"\b(?:Figure|Table|Section|Chapter) (?:\d+(?:\.\d+)*|[A-Z](?:\.\d+)*)\b")
INLINE = re.compile(
    r"(\$[^$\n]+\$|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`|"
    r"\[[^]]+\]\([^)]+\)|\b(?:Figure|Table|Section|Chapter) "
    r"(?:\d+(?:\.\d+)*|[A-Z](?:\.\d+)*)\b)"
)


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def key(label: str) -> str:
    return label.lower().replace(" ", "_").replace(".", "_")


def field(paragraph, instruction: str, cached: str = "") -> None:
    for tag, value in (
        ("begin", None),
        ("instruction", instruction),
        ("separate", None),
        ("text", cached),
        ("end", None),
    ):
        run = OxmlElement("w:r")
        if tag in {"begin", "separate", "end"}:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        elif tag == "instruction":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = " " + value + " "
        else:
            element = OxmlElement("w:t")
            element.text = value
        run.append(element)
        paragraph._p.append(run)


class ThesisWordBuilder:
    def __init__(self, template: Path, layout: dict, math_manifest: dict, stylesheet: Path):
        if sha(template) != TEMPLATE_SHA:
            raise ValueError("The reviewed Word template changed; review it before rebuilding")
        self.document = Document(template)
        self.layout = layout
        self.math = math_manifest
        self.transform = etree.XSLT(etree.parse(str(stylesheet)))
        self.bookmark_id = 10000
        self.math_count = 0
        self.figures = []
        self.tables = []
        self.in_abstract = False
        self.in_references = False
        self.cover = json.loads((ROOT / "docs/academic-cover.json").read_text(encoding="utf-8"))
        self.width = (
            self.document.sections[0].page_width
            - self.document.sections[0].left_margin
            - self.document.sections[0].right_margin
        )
        self._prepare_template()

    def _prepare_template(self):
        document = self.document
        declaration = next(p for p in document.paragraphs if p.text.strip() == "DECLARATION")
        body = document._element.body
        index = list(body).index(declaration._p)
        for element in list(body)[index:]:
            if element.tag != qn("w:sectPr"):
                body.remove(element)
        for paragraph in document.paragraphs:
            if paragraph.text == self.cover["project_title"]:
                paragraph.style = document.styles["Title"]
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(15)
                    run.bold = True
                    run.font.color.rgb = RGBColor(0, 0, 0)
            if paragraph.text.startswith("Supervisor:"):
                paragraph.text = "Supervisor: " + self.cover["supervisor"]
                paragraph.runs[0].bold = True
        for name in ("Heading 1", "Heading 2", "Heading 3"):
            style = document.styles[name]
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.keep_with_next = True
            if name == "Heading 1":
                style.paragraph_format.page_break_before = True
        for level in (1, 2):
            style = document.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            style.style_id = f"TOC{level}"
            style.base_style = document.styles["Normal"]
            style.font.size = Pt(11)
            style.paragraph_format.line_spacing = Pt(13)
            style.paragraph_format.space_after = Pt(2)
            if level == 2:
                style.paragraph_format.left_indent = Pt(12)
        for name, size, bold in (
            ("Figure Caption", 10, False),
            ("Table Caption", 10, True),
            ("Frontmatter Heading", 16, True),
            ("Contents Heading", 16, True),
            ("Arabic Thesis", 11.5, False),
        ):
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = document.styles["Normal"]
            style.font.name = "Noto Naskh Arabic" if name == "Arabic Thesis" else "Times New Roman"
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.keep_with_next = name == "Table Caption"
            if name in {"Frontmatter Heading", "Contents Heading"}:
                style.paragraph_format.page_break_before = True
                style.paragraph_format.space_after = Pt(20)
                outline = OxmlElement("w:outlineLvl")
                outline.set(qn("w:val"), "0" if name == "Frontmatter Heading" else "9")
                style.element.get_or_add_pPr().append(outline)
        settings = document.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
        document.core_properties.author = "; ".join(a["name"] for a in self.cover["authors"])
        document.core_properties.title = self.cover["project_title"]
        document.core_properties.subject = "Graduation project thesis"
        document.core_properties.comments = ""

    def bookmark(self, paragraph, label: str):
        self.bookmark_id += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_id))
        start.set(qn("w:name"), key(label))
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self.bookmark_id))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def equation(self, paragraph, latex: str, display: bool = False):
        identifier = hashlib.sha256(f"{display}:{latex.strip()}".encode()).hexdigest()[:20]
        item = self.math[identifier]
        mathml = etree.parse(item["mathml"])
        namespace = "http://www.w3.org/1998/Math/MathML"
        operators = []
        for operator in mathml.findall(".//{" + namespace + "}mo"):
            if operator.text and len(operator.text) > 1 and operator.text.isalpha():
                # Office's XSLT mistakes named operators inside fences for separators.
                # Keep the complete upright identifier instead of truncating it.
                operators.append(operator.text)
                operator.tag = "{" + namespace + "}mi"
                operator.set("mathvariant", "normal")
        transformed = self.transform(mathml)
        node = transformed.getroot()
        plain = "".join(node.itertext())
        if any(operator not in plain for operator in operators):
            raise ValueError(f"MathML conversion lost an operator: {latex}")
        if not display and node.tag.endswith("oMathPara"):
            node = node.find("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        paragraph._p.append(copy.deepcopy(node))
        self.math_count += 1

    def reference(self, paragraph, label: str):
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), key(label))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def inline(self, paragraph, text: str, *, references: bool = True):
        for token in INLINE.split(text):
            if not token:
                continue
            if token.startswith("$") and token.endswith("$"):
                self.equation(paragraph, token[1:-1])
            elif token.startswith("**"):
                paragraph.add_run(token[2:-2]).bold = True
            elif token.startswith("*"):
                paragraph.add_run(token[1:-1]).italic = True
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            elif token.startswith("["):
                match = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token)
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(
                    qn("r:id"),
                    self.document.part.relate_to(match[2], RT.HYPERLINK, is_external=True),
                )
                run = OxmlElement("w:r")
                text_node = OxmlElement("w:t")
                text_node.text = match[1]
                run.append(text_node)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            elif references and REFERENCE.fullmatch(token):
                self.reference(paragraph, token)
            else:
                paragraph.add_run(token)

    def paragraph(self, text: str, style: str | None = None):
        arabic = bool(re.search(r"[\u0600-\u06ff]", text))
        paragraph = self.document.add_paragraph(style="Arabic Thesis" if arabic else style)
        paragraph.paragraph_format.widow_control = True
        if self.in_references:
            paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(8)
        if arabic:
            properties = paragraph._p.get_or_add_pPr()
            properties.append(OxmlElement("w:bidi"))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = Pt(14.5)
            paragraph.paragraph_format.space_after = Pt(6)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.inline(paragraph, text)
        if self.in_abstract and not arabic:
            paragraph.paragraph_format.line_spacing = Pt(14)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                run.font.size = Pt(11)
        if arabic:
            for run in paragraph.runs:
                run.font.size = Pt(11.5)
                properties = run._r.get_or_add_rPr()
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:cs"), "Noto Naskh Arabic")
                properties.append(fonts)
                size = OxmlElement("w:szCs")
                size.set(qn("w:val"), "23")
                properties.append(size)
                rtl = OxmlElement("w:rtl")
                properties.append(rtl)
        return paragraph

    def heading(self, level: int, title: str):
        if level == 1:
            self.in_abstract = title in {"Abstract", "الملخص"}
            self.in_references = title == "References"
        paragraph = self.document.add_paragraph(style="Heading " + str(min(level, 3)))
        if re.search(r"[\u0600-\u06ff]", title):
            paragraph.add_run(title).font.name = "Noto Naskh Arabic"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        else:
            self.inline(paragraph, title, references=False)
        number = re.match(r"(\d+(?:\.\d+)*|[A-Z]\.\d+)\s", title)
        if number:
            self.bookmark(paragraph, ("Chapter " if level == 1 else "Section ") + number[1])
        return paragraph

    def frontmatter(self):
        for title, instruction in (
            ("Contents", 'TOC \\o "1-2" \\t "Frontmatter Heading,1" \\h \\z \\u'),
            ("List of figures", "TOC \\f F \\h \\z"),
            ("List of tables", 'TOC \\t "Table Caption,1" \\h \\z'),
        ):
            style = "Contents Heading" if title == "Contents" else "Frontmatter Heading"
            self.document.add_paragraph(title, style=style)
            paragraph = self.document.add_paragraph()
            field(paragraph, instruction, "Updated by Microsoft Word during verification.")
        self.document.add_paragraph("List of abbreviations", style="Frontmatter Heading")
        self.table([["Abbreviation", "Meaning"], *self.layout["abbreviations"]], count=False)

    def table(self, rows: list[list[str]], *, count: bool = True):
        table = self.document.add_table(rows=1, cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        weights = [min(40, max(12, max(len(row[i]) for row in rows))) for i in range(len(rows[0]))]
        widths = [int(self.width * value / sum(weights)) for value in weights]
        for column, width in zip(table.columns, widths, strict=True):
            column.width = width
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for i, row in enumerate(rows):
            cells = table.rows[0].cells if i == 0 else table.add_row().cells
            for j, value in enumerate(row):
                cell = cells[j]
                cell.width = widths[j]
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    line = OxmlElement("w:" + edge)
                    line.set(qn("w:val"), "single")
                    line.set(qn("w:sz"), "4")
                    line.set(qn("w:color"), "D9D9D9")
                    borders.append(line)
                properties.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "173F43" if i == 0 else ("F4F6F5" if i % 2 else "FFFFFF"))
                properties.append(shading)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.keep_with_next = i == 0 or (
                    (len(rows) <= 10 or (len(rows) <= 16 and len(rows[0]) == 2))
                    and i < len(rows) - 1
                )
                self.inline(paragraph, value)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            table.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        self.document.add_paragraph().paragraph_format.space_after = Pt(2)
        if count:
            self.tables.append(rows[0])

    def figure(self, name: str):
        from PIL import Image

        item = self.layout["figure_metadata"][name]
        path = ROOT / item["path"]
        with Image.open(path) as image:
            ratio = image.height / image.width
        width = min(self.width, Inches(6.7) / ratio)
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        picture = paragraph.add_run().add_picture(str(path), width=int(width))
        picture._inline.docPr.set("descr", item["title"] + ". " + item["explanation"])
        paragraph = self.document.add_paragraph(style="Figure Caption")
        label = "Figure " + item["number"]
        self.inline(
            paragraph, label + ". " + item["title"] + ". " + item["explanation"], references=False
        )
        self.bookmark(paragraph, label)
        field(paragraph, f'TC "{label}. {item["title"]}" \\f F \\l 1')
        self.figures.append(label)

    def body(self, markdown: str):
        lines = markdown.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index].strip()
            if not line:
                index += 1
                continue
            heading = re.match(r"^(#{1,3}) (.+)$", line)
            if heading:
                if line.startswith("# 1 "):
                    self.frontmatter()
                self.heading(len(heading[1]), heading[2])
                index += 1
                continue
            if line.startswith("$$"):
                if self.document.paragraphs and self.document.paragraphs[-1].text.strip():
                    self.document.paragraphs[-1].paragraph_format.keep_with_next = True
                formula = [line[2:]]
                index += 1
                while "$$" not in formula[-1]:
                    formula.append(lines[index])
                    index += 1
                latex = "\n".join(formula).rsplit("$$", 1)[0].strip()
                paragraph = self.document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.space_after = Pt(10)
                self.equation(paragraph, latex, display=True)
                continue
            marker = re.fullmatch(r"\[\[FIGURE:([a-z]+)\]\]", line)
            if marker:
                self.figure(marker[1])
                index += 1
                continue
            caption = re.fullmatch(r"\*\*(Table \d+)\. (.+)\*\*", line)
            if caption:
                paragraph = self.document.add_paragraph(style="Table Caption")
                self.inline(paragraph, caption[1] + ". " + caption[2], references=False)
                self.bookmark(paragraph, caption[1])
                index += 1
                continue
            if line.startswith("|"):
                rows = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    row = lines[index].strip()
                    if not re.fullmatch(r"\|?[\s|:-]+\|?", row):
                        rows.append([cell.strip() for cell in row.strip("|").split("|")])
                    index += 1
                self.table(rows)
                continue
            if line.startswith("```"):
                index += 1
                while index < len(lines) and not lines[index].startswith("```"):
                    paragraph = self.document.add_paragraph()
                    run = paragraph.add_run(lines[index])
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                    paragraph.paragraph_format.space_after = Pt(0)
                    index += 1
                index += 1
                continue
            paragraph_lines = [line]
            index += 1
            while index < len(lines) and lines[index].strip():
                candidate = lines[index].strip()
                if candidate.startswith(("#", "|", "$$", "[[FIGURE:", "**Table ", "```")):
                    break
                if re.match(r"(?:[-*]|\d+[.)])\s+", candidate):
                    break
                paragraph_lines.append(candidate)
                index += 1
            text = " ".join(paragraph_lines)
            self.paragraph(text)


def build(
    template: Path, source: Path, layout_path: Path, math_path: Path, stylesheet: Path, output: Path
):
    layout = json.loads(layout_path.read_text(encoding="utf-8"))
    if sha(source) != layout["resolved_source_sha256"]:
        raise ValueError("Resolved Word source differs from the reviewed PDF build")
    builder = ThesisWordBuilder(template, layout, json.loads(math_path.read_text()), stylesheet)
    builder.body(source.read_text(encoding="utf-8"))
    output.parent.mkdir(parents=True, exist_ok=True)
    builder.document.save(output)
    receipt = {
        "schema": "gridguard.focused_thesis_word.1",
        "docx_sha256": sha(output),
        "template_sha256": sha(template),
        "source_sha256": sha(source),
        "pdf_build_sha256": sha(layout_path),
        "builder_sha256": sha(Path(__file__)),
        "cover_metadata_sha256": sha(ROOT / "docs/academic-cover.json"),
        "office_math_transform_sha256": sha(stylesheet),
        "native_math_count": builder.math_count,
        "figures": builder.figures,
        "numbered_table_count": len(builder.tables),
        "fields_require_refresh": True,
    }
    output.with_name(output.name + ".build.json").write_text(
        json.dumps(receipt, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--template", type=Path, default=ROOT / "docs/source/reviewed-thesis-template.docx"
    )
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--layout", required=True, type=Path)
    parser.add_argument(
        "--math-manifest", type=Path, default=ROOT / "tmp/math-typesetting/rendered.json"
    )
    parser.add_argument(
        "--stylesheet",
        type=Path,
        default=Path("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "output/documents/updated/gridguard-khartoum-thesis.docx",
    )
    args = parser.parse_args()
    build(args.template, args.source, args.layout, args.math_manifest, args.stylesheet, args.output)
