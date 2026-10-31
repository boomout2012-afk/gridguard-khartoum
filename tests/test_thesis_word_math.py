"""Guard against native Word math conversion silently dropping operator letters."""

import hashlib
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from tools.build_focused_thesis_docx import ThesisWordBuilder


def builder_for_math(tmp_path: Path, *, preserve_identifiers: bool) -> ThesisWordBuilder:
    mathml = tmp_path / "minimum.mml"
    mathml.write_text(
        '<math xmlns="http://www.w3.org/1998/Math/MathML">'
        "<mrow><mo>(</mo><mo>min</mo><mi>x</mi><mo>)</mo></mrow></math>",
        encoding="utf-8",
    )
    value = "." if preserve_identifiers else "substring(., 1, 1)"
    transform = etree.XSLT(
        etree.fromstring(
            f'''<xsl:stylesheet version="1.0"
                xmlns:xsl="http://www.w3.org/1999/XSL/Transform"
                xmlns:mml="http://www.w3.org/1998/Math/MathML"
                xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
                <xsl:template match="/">
                  <m:oMath><m:r><m:t>
                    <xsl:for-each select="//mml:mi">
                      <xsl:value-of select="{value}"/>
                    </xsl:for-each>
                  </m:t></m:r></m:oMath>
                </xsl:template>
              </xsl:stylesheet>'''.encode()
        )
    )
    builder = ThesisWordBuilder.__new__(ThesisWordBuilder)
    identifier = hashlib.sha256(b"False:\\min(x)").hexdigest()[:20]
    builder.math = {identifier: {"mathml": str(mathml)}}
    builder.transform = transform
    builder.math_count = 0
    return builder


def test_named_operator_becomes_complete_upright_native_math(tmp_path: Path) -> None:
    builder = builder_for_math(tmp_path, preserve_identifiers=True)
    paragraph = Document().add_paragraph()
    builder.equation(paragraph, r"\min(x)")
    nodes = paragraph._p.findall(".//{http://schemas.openxmlformats.org/officeDocument/2006/math}t")
    assert "".join(node.text for node in nodes) == "minx"
    assert builder.math_count == 1


def test_native_math_conversion_fails_if_operator_is_truncated(tmp_path: Path) -> None:
    builder = builder_for_math(tmp_path, preserve_identifiers=False)
    with pytest.raises(ValueError, match="lost an operator"):
        builder.equation(Document().add_paragraph(), r"\min(x)")
    assert builder.math_count == 0
